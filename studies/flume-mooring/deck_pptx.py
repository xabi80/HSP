"""PowerPoint copy of a Slides deck folder (deck.json + slides/*.html, as leadership-deck/ and
mooring-design-study/): one full-slide picture per slide, rendered with headless Edge from the same
slide HTML, stylesheet and images as the folder's render_pdf.py, and each slide's speaker notes
(its <aside>) in the PowerPoint notes. The slide text is a picture, not editable: the editable
export is the online deck's Share > Export > PowerPoint.

Writes <deck folder>/<the PDF's name>.pptx, e.g. mooring-design-study/mooring-design-study.pptx.
Run: python deck_pptx.py mooring-design-study   (or leadership-deck)
"""

from __future__ import annotations

import base64
import html
import importlib.util
import json
import re
import subprocess
import sys
import time
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
W_PX, H_PX = 1920, 1080


def _render_module(deck_dir: Path):  # type: ignore[no-untyped-def]
    spec = importlib.util.spec_from_file_location("render_pdf", deck_dir / "render_pdf.py")
    rp = importlib.util.module_from_spec(spec)  # type: ignore[arg-type]
    spec.loader.exec_module(rp)  # type: ignore[union-attr]
    return rp


def _image_uri(rp, value) -> str:  # type: ignore[no-untyped-def]
    """IMAGES maps an asset id to a figure name (trimmed by rp.trim) or to a file path."""
    data = rp.trim(value) if isinstance(value, str) else Path(value).read_bytes()
    return "data:image/png;base64," + base64.b64encode(data).decode()


def _shot(edge: str, page: Path, png: Path) -> None:
    png.unlink(missing_ok=True)
    subprocess.run([edge, "--headless=new", "--disable-gpu", "--hide-scrollbars",
                    f"--window-size={W_PX},{H_PX}", "--virtual-time-budget=8000",
                    f"--screenshot={png}", page.as_uri()], capture_output=True, timeout=120)
    last = -1  # Edge may finish writing after it returns
    for _ in range(60):
        size = png.stat().st_size if png.exists() else -1
        if size > 0 and size == last:
            return
        last = size
        time.sleep(0.5)
    raise RuntimeError(f"Edge did not write {png.name}")


def main() -> None:
    deck_dir = (HERE / sys.argv[1]).resolve()
    rp = _render_module(deck_dir)
    deck = json.loads((deck_dir / "deck.json").read_text(encoding="utf-8"))
    pdfs = sorted(deck_dir.glob("*.pdf"))
    if len(pdfs) != 1:
        raise ValueError(f"{deck_dir.name}: expected one PDF to name the .pptx, found {len(pdfs)}")
    links = "".join(f'<link rel="stylesheet" href="{html.escape(f["href"])}">'
                    for f in deck["faces"].values() if "href" in f)
    head = (f"<!doctype html><html><head><meta charset='utf-8'>{links}<style>{rp.CSS}</style>"
            "</head><body>")
    uris = {aid: _image_uri(rp, v) for aid, v in rp.IMAGES.items()}

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for sid in deck["order"]:
        s = (deck_dir / "slides" / f"{sid}.html").read_text(encoding="utf-8")
        for aid, uri in uris.items():
            s = s.replace(f"/_blob/{aid}", uri)
        page, png = deck_dir / f"_pptx_{sid}.html", deck_dir / f"_pptx_{sid}.png"
        page.write_text(head + s + "</body></html>", encoding="utf-8")
        try:
            _shot(str(rp.EDGE), page, png)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
        finally:
            page.unlink(missing_ok=True)
            png.unlink(missing_ok=True)
        m = re.search(r"<aside>(.*?)</aside>", s, re.S)
        if m:
            slide.notes_slide.notes_text_frame.text = html.unescape(m.group(1).strip())
    out = deck_dir / (pdfs[0].stem + ".pptx")
    prs.save(str(out))
    print(f"wrote {out.relative_to(HERE)} ({len(deck['order'])} slides, notes included)")


if __name__ == "__main__":
    main()
