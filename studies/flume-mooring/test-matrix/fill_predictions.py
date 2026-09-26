"""Fill the HSFP 1:50 test matrix with FloatSim's pre-test predictions.

Reads the original workbook (HSFP_1-50_Wave_Test_Matrix_9-26-2026.xlsx, kept unchanged next to
this script) and writes *_FloatSim.xlsx:

- Regular_Waves rows 24-26 (the natural-period placeholders): the predicted heave and tilt
  natural periods at full scale, with notes. Their Full/Reduced flags stay "N", so the run
  counts and tank time are unchanged until the matrix owner switches them on.
- A new sheet, FloatSim_Predictions:
  - the heave and tilt natural periods per article, free and moored, on the flume BEM and
    corrected to the converged waterline mesh (DESIGN-BASIS C8);
  - the FloatSim damping and half-power bands;
  - the moored slow-mode periods, with the decay record each needs;
  - the recommended fine-step tilt band per configuration.

Every number comes from the study's records:
- attachment_sweep (the chosen operational designs);
- extreme_set.json (declare);
- bem_waterline.json (the mesh correction);
- resonance_bandwidth.json (damping);
- spec_statics*.json (slow modes).

Full-scale periods are formulas on the workbook's own scale factor (Inputs!B4). Recalculate in
Excel afterwards: the build does it through Excel's COM interface (see the README).

Run: python fill_predictions.py
"""

# ruff: noqa: E402  -- sys.path bootstrap for the study modules precedes their import
from __future__ import annotations

import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
STUDY = HERE.parent
sys.path.insert(0, str(STUDY))

import attachment_sweep as asw
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = HERE / "HSFP_1-50_Wave_Test_Matrix_9-26-2026.xlsx"
OUT = HERE / "HSFP_1-50_Wave_Test_Matrix_9-26-2026_FloatSim.xlsx"
ARTS = ("buoy", "cluster", "platform")
NAME = {"buoy": "Single buoy", "cluster": "Cluster (4 buoys)", "platform": "4x4 platform"}
BAND_HALF, STEP = 0.20, 0.05  # fine band: centre ± 0.20 s model, 0.05 s model steps


def load(name: str) -> dict:
    return json.loads((STUDY / name).read_text(encoding="utf-8"))


def predictions() -> dict:
    """Natural periods (model s) per article: flume BEM, and the converged-mesh estimate."""
    bw = load("bem_waterline.json")["articles"]
    rb = load("resonance_bandwidth.json")["cases"]
    ext = load("extreme_set.json")["declare"]
    shift = {a: bw[a]["12"]["shift_vs_nt96_s"] for a in ARTS}
    shift_free_buoy = bw["buoy_free"]["12"]["shift_vs_nt96_s"]
    rows = []
    for a in ARTS:
        r = asw.chosen(a)["row"]
        moored = {"heave": r["heave_T_s"], "tilt": r["tilt_T_s"]}
        free = {m: moored[m] / (1 + r[f"{m}_shift_pct"] / 100) for m in ("heave", "tilt")}
        rbc = rb[f"{a}@H0.04"]
        for mode in ("heave", "tilt"):
            zeta = rbc["pitch" if mode == "tilt" else "heave"]["zeta"]
            band = rbc["pitch" if mode == "tilt" else "heave"]["dT_hp"]
            label = "pitch" if (a == "buoy" and mode == "tilt") else mode
            variants = [("free", free[mode]), ("moored, operational cords", moored[mode])]
            if a in ext:
                variants.append(("moored, extreme cords", ext[a][f"{mode}_T_s"]))
            for moor, t in variants:
                if a == "buoy" and moor == "free":
                    d, basis = shift_free_buoy[mode], "converged value computed (DESIGN-BASIS C8)"
                else:
                    d = shift[a][mode]
                    basis = (
                        "converged value estimated: the C8 mesh shift of this article's earlier "
                        "moored design"
                    )
                rows.append(
                    {
                        "article": NAME[a],
                        "mode": label,
                        "mooring": moor,
                        "flume": t,
                        "converged": t - d,
                        "zeta": zeta,
                        "band": band,
                        "basis": basis,
                    }
                )
    return {"rows": rows}


def main() -> None:
    wb = openpyxl.load_workbook(SRC)
    lam = float(wb["Inputs"]["B4"].value)
    sq = np.sqrt(lam)
    pr = predictions()["rows"]

    def conv(article: str, mode: str, moor: str) -> float:
        return next(
            r["converged"]
            for r in pr
            if r["article"] == NAME[article] and r["mode"] == mode and r["mooring"] == moor
        )

    op = "moored, operational cords"
    heave = {a: conv(a, "heave", op) for a in ARTS}
    tilt = {a: conv(a, "pitch" if a == "buoy" else "tilt", op) for a in ARTS}

    # ---- Regular_Waves placeholder rows 24-26 ----
    ws = wb["Regular_Waves"]
    fs = lambda t: float(round(float(t) * sq, 1))  # noqa: E731  -- local one-line conversion
    fills = {
        24: (
            fs(np.mean(list(heave.values()))),
            "FloatSim pre-test: heave natural period, moored: buoy {:.1f}, cluster {:.1f}, "
            "platform {:.1f} s FS. Broad peak (zeta ~11 %): the 1 s grid resolves it. The heave "
            "RAO peaks later in waves (~18.5-19 s FS). Replace with the measured decay. See "
            "FloatSim_Predictions.".format(*(heave[a] * sq for a in ARTS)),
        ),
        25: (
            fs(tilt["cluster"]),
            "FloatSim pre-test: tilt/pitch natural period, moored (operational cords): buoy "
            "{:.1f}, cluster {:.1f}, platform {:.1f} s FS. Sharp peak (zeta ~3-4 %): needs the "
            "fine-step band per configuration (FloatSim_Predictions). Replace with the measured "
            "decay.".format(*(tilt[a] * sq for a in ARTS)),
        ),
        26: (
            fs(tilt["platform"]) + 0.5,
            "FloatSim pre-test: platform tilt {:.1f} s FS + 0.5 s, the resonance's upper flank, "
            "beyond the 20 s bound. Extend the sweep to ~22 s FS (FloatSim_Predictions).".format(
                tilt["platform"] * sq
            ),
        ),
    }
    for r, (val, note) in fills.items():
        ws[f"A{r}"] = val
        f = ws[f"T{r}"].value
        old = f.split('&"')[-1].rstrip('")')  # the placeholder label, as the formula ends with it
        assert f.count(old) == 2, (r, old)
        ws[f"T{r}"] = f.replace(old, note.replace('"', "'"))

    # ---- FloatSim_Predictions sheet ----
    sh = wb.create_sheet("FloatSim_Predictions")
    bold, head = Font(bold=True), PatternFill("solid", fgColor="D9E1E2")
    wrap = Alignment(wrap_text=True, vertical="top")
    rows: list[list] = [
        ["FloatSim pre-test predictions - flume mooring study (HSP repo, studies/flume-mooring)"],
        [
            "Pre-test predictions, to be replaced by the measured free decays. Model scale in "
            "fresh water; full scale = model x √λ (Inputs!B4). FloatSim LEVEL1, linear "
            "hydrodynamics, Morison drag. The flume BEM's 12-sided waterline makes periods long; "
            "converged values per DESIGN-BASIS C8. Pitch inputs (CoG, inertia) are unmeasured: "
            "±3 % (DESIGN-BASIS Phase I)."
        ],
        [],
        ["1. Natural periods (heave, tilt/pitch)"],
        [
            "Article",
            "Mode",
            "Mooring",
            "Flume BEM T (model s)",
            "Converged T (model s)",
            "Converged T (full scale s)",
            "ζ (FloatSim, H = 0.04 m)",
            "Half-power band (model s)",
            "Basis",
        ],
    ]
    for r in pr:
        rows.append(
            [
                r["article"],
                r["mode"],
                r["mooring"],
                round(r["flume"], 3),
                round(r["converged"], 3),
                None,
                round(r["zeta"], 3),
                round(r["band"], 3),
                r["basis"],
            ]
        )
    first_nat = 6
    rows += [
        [
            "ζ and the half-power band: FloatSim regular-wave sweeps at H = 0.04 m model with the "
            "earlier (pin-level) mooring (resonance_bandwidth.json). They grow with wave height "
            "(cluster tilt ζ 6.4 % at H = 0.12 m)."
        ],
        [
            "In waves the heave RAO peaks later than the decay period: ~2.65-2.70 s model on the "
            "flume BEM (≈ 18.5-19 s FS after the mesh correction)."
        ],
        [],
        ["2. Moored slow modes (MOORING-SPEC rev C.1 §5) and the decay record they need"],
        [
            "Article",
            "Cord set",
            "Surge T (model s)",
            "Sway T (model s)",
            "Yaw T (model s)",
            "Surge T (full scale s)",
            "Sway T (full scale s)",
            "Record for 5 cycles of the slowest " "soft mode (model min)",
            "Note",
        ],
    ]
    st, sx = load("spec_statics.json"), load("spec_statics_extreme.json")
    first_slow = len(rows) + 1
    for a in ARTS:
        for tag, src in (("operational", st), ("extreme", sx)):
            if a not in src:
                continue
            p = src[a]["periods_s"]
            soft = max(p["surge"], p["sway"], p["yaw"] if p["yaw"] > 5 else 0.0)
            rows.append(
                [
                    NAME[a],
                    "one set" if a == "buoy" else tag,
                    round(p["surge"], 2),
                    round(p["sway"], 2),
                    round(p["yaw"], 2),
                    None,
                    None,
                    round(5 * soft / 60, 1),
                    "yaw is stiff (collar)" if p["yaw"] < 5 else "",
                ]
            )
    last_slow = len(rows)
    rows += [
        [
            "The matrix's free-decay run is 2 min (Inputs!B21). The moored sway and yaw decays "
            "need up to ~5 min. The slow modes are also where irregular seas drive slow drift, "
            "which FloatSim does not model."
        ],
        [],
        [
            "3. Recommended fine-step tilt band per configuration (centre = converged moored tilt, "
            f"± {BAND_HALF:.2f} s model, {STEP:.2f} s model steps; ≥ 95 % of the peak)"
        ],
        [
            "Configuration",
            "Centre T (model s)",
            "Band periods (model s)",
            None,
            None,
            "Band periods (full scale s)",
            None,
            None,
            "Note",
        ],
    ]
    first_band = len(rows) + 1
    for a in ARTS:
        c = tilt[a]
        ts = np.round(np.arange(c - BAND_HALF, c + BAND_HALF + 1e-9, STEP), 3)
        rows.append(
            [
                NAME[a] + ", operational cords",
                round(c, 3),
                ", ".join(f"{t:.2f}" for t in ts),
                None,
                None,
                ", ".join(f"{t * sq:.1f}" for t in ts),
                None,
                None,
                "Re-centre on the measured free decay. Heights: add constant small levels "
                "(0.02 and 0.04 m model) to stay near-linear at resonance.",
            ]
        )
    for row in rows:
        sh.append(row)
    for r in range(first_nat, first_nat + len(pr)):  # full-scale formulas
        sh[f"F{r}"] = f"=E{r}*SQRT(Inputs!$B$4)"
    for r in range(first_slow, last_slow + 1):
        sh[f"F{r}"] = f"=C{r}*SQRT(Inputs!$B$4)"
        sh[f"G{r}"] = f"=D{r}*SQRT(Inputs!$B$4)"
    sh["A1"].font = Font(bold=True, size=13)
    for r in (4, first_slow - 2, first_band - 2):
        sh[f"A{r}"].font = bold
    for r in (5, first_slow - 1, first_band - 1):
        for c in range(1, 10):
            sh.cell(r, c).font, sh.cell(r, c).fill = bold, head
    for c, w in enumerate((24, 10, 26, 16, 16, 16, 14, 18, 60), start=1):
        sh.column_dimensions[get_column_letter(c)].width = w
    for row in sh.iter_rows():
        for cell in row:
            cell.alignment = wrap
    for r in range(first_nat, first_nat + len(pr)):
        for c in "DEFGH":
            sh[f"{c}{r}"].number_format = "0.000" if c in "DEH" else ("0.0" if c == "F" else "0.0%")
    for r in range(first_slow, last_slow + 1):
        for c in "FG":
            sh[f"{c}{r}"].number_format = "0"
    wb.save(OUT)
    print(f"wrote {OUT.name}: rows 24-26 = {[v for v, _ in fills.values()]} s FS")


if __name__ == "__main__":
    main()
