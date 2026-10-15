"""Shared slide helpers for the payload-stability decks (make_payload_ppt.py, make_platform_ppt.py):
the palette, one 13.333 x 7.5 in presentation, text / card / stat / picture helpers and native
PowerPoint XY charts. Each deck script runs in its own process and saves `prs`.
"""

# ruff: noqa: E501, E702  -- compact slide-assembly lines, as the deck scripts
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent

NAVY, INK, MUTED = RGBColor(0x14, 0x23, 0x2E), RGBColor(0x2B, 0x3A, 0x44), RGBColor(0x5C, 0x6B, 0x75)
TEAL, WATER = RGBColor(0x0C, 0x8B, 0x96), RGBColor(0xE4, 0xF2, 0xF4)
AMBER, RED = RGBColor(0xC9, 0x77, 0x2B), RGBColor(0xB2, 0x43, 0x2C)
GREEN, WHITE = RGBColor(0x5A, 0x8A, 0x2B), RGBColor(0xFF, 0xFF, 0xFF)
ICE = RGBColor(0xB9, 0xD6, 0xDC)
DIR_COL = {0.0: TEAL, 15.0: GREEN, 30.0: AMBER, 45.0: RED}
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(dark: bool = False, notes: str = ""):  # type: ignore[no-untyped-def]
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY if dark else WHITE
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def text(s, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,  # type: ignore[no-untyped-def]
         anchor=MSO_ANCHOR.TOP, gap=6):
    """paras: a str, or a list of paragraphs; a paragraph is a str or a list of
    (text, {bold, color, size, italic}) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.0)
    for i, para in enumerate([paras] if isinstance(paras, str) else paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(gap)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run(); r.text = t
            r.font.name = FONT; r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold); r.font.italic = o.get("italic", False)
            r.font.color.rgb = o.get("color", color)
    return tb


def title(s, t, sub=None, dark=False):  # type: ignore[no-untyped-def]
    text(s, 0.6, 0.38, 12.1, 0.75, t, size=32, bold=True, color=WHITE if dark else NAVY)
    if sub:
        text(s, 0.6, 1.08, 12.1, 0.45, sub, size=16, color=ICE if dark else MUTED)


def card(s, x, y, w, h, fill=WATER):  # type: ignore[no-untyped-def]
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.06
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def stat(s, x, y, w, big, small, color=TEAL, fill=WATER, h=1.65):  # type: ignore[no-untyped-def]
    card(s, x, y, w, h, fill)
    text(s, x + 0.25, y + 0.15, w - 0.5, 0.75, big, size=34, bold=True, color=color)
    text(s, x + 0.25, y + 0.88, w - 0.5, h - 0.95, small, size=13, color=INK)


def picture(s, path, x, y, w=None, h=None):  # type: ignore[no-untyped-def]
    return s.shapes.add_picture(str(HERE / path), Inches(x), Inches(y),
                                Inches(w) if w else None, Inches(h) if h else None)


def fit_picture(s, path, x, y, w, h):  # type: ignore[no-untyped-def]
    """The picture scaled to fit inside the box (x, y, w, h), centred in it."""
    from PIL import Image
    iw, ih = Image.open(path if Path(path).is_absolute() else HERE / path).size
    k = min(w / iw, h / ih)
    pw, ph = iw * k, ih * k
    return s.shapes.add_picture(str(path if Path(path).is_absolute() else HERE / path),
                                Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw))


def xy_chart(s, x, y, w, h, series, xt, yt, xr=None, yr=None, legend=True,  # type: ignore[no-untyped-def]
             title_text=None):
    """series: list of dicts {name, x, y, color, width, dash, marker}."""
    cd = XyChartData()
    for sr in series:
        se = cd.add_series(sr["name"])
        for a, b in zip(sr["x"], sr["y"], strict=True):
            se.add_data_point(float(a), float(b))
    gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(x), Inches(y),
                            Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.name = FONT; ch.font.size = Pt(11); ch.font.color.rgb = MUTED
    ch.has_title = title_text is not None
    if title_text:
        ch.chart_title.text_frame.text = title_text
        tp = ch.chart_title.text_frame.paragraphs[0]
        tp.runs[0].font.size = Pt(13); tp.runs[0].font.bold = True; tp.runs[0].font.color.rgb = INK
    for sr, plot_s in zip(series, ch.plots[0].series, strict=True):
        ln = plot_s.format.line
        ln.color.rgb = sr["color"]; ln.width = Pt(sr.get("width", 2.25))
        if sr.get("dash"):
            ln.dash_style = sr["dash"]
        plot_s.smooth = False
        if sr.get("marker"):
            plot_s.marker.style = sr["marker"]; plot_s.marker.size = 6
            plot_s.marker.format.fill.solid(); plot_s.marker.format.fill.fore_color.rgb = sr["color"]
            plot_s.marker.format.line.color.rgb = sr["color"]
        else:
            plot_s.marker.style = XL_MARKER_STYLE.NONE
    for ax, t, rng in ((ch.category_axis, xt, xr), (ch.value_axis, yt, yr)):
        ax.has_title = True; ax.axis_title.text_frame.text = t
        r0 = ax.axis_title.text_frame.paragraphs[0].runs[0]
        r0.font.size = Pt(12); r0.font.bold = False; r0.font.color.rgb = INK
        ax.tick_labels.font.size = Pt(11); ax.tick_labels.font.color.rgb = MUTED
        ax.format.line.color.rgb = ICE
        ax.has_major_gridlines = True; ax.major_gridlines.format.line.color.rgb = RGBColor(0xE6, 0xEC, 0xEE)
        if rng:
            ax.minimum_scale, ax.maximum_scale = rng
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(11); ch.legend.font.color.rgb = INK
    return ch


def arrow(s, x, y, w=0.45, h=0.32, color=TEAL):  # type: ignore[no-untyped-def]
    sh = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color; sh.line.fill.background()


def table(s, x, y, w, h, cells, size=13):  # type: ignore[no-untyped-def]
    """A styled table: cells[0] is the header row (teal), then banded rows."""
    n, k = len(cells), len(cells[0])
    tb = s.shapes.add_table(n, k, Inches(x), Inches(y), Inches(w), Inches(h)).table
    for i, row_ in enumerate(cells):
        for j, v in enumerate(row_):
            c = tb.cell(i, j); c.text = v
            c.fill.solid(); c.fill.fore_color.rgb = TEAL if i == 0 else (WATER if i % 2 else WHITE)
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size); r.font.name = FONT
                    r.font.color.rgb = WHITE if i == 0 else INK; r.font.bold = i == 0
    return tb
