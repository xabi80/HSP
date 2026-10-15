"""PowerPoint deck of the cluster payload-stability study: pictures, native charts and summary.

Every number comes from this folder's records:
- buoy_hydrostatics_vs_draft.json;
- cluster_envelope.json;
- tilt_paths.json;
- floatsim_crosscheck.json.
The pictures come from payload_figures.py: sketch_cluster.png, seq_*.png and the anim_*.gif
animations, which play in PowerPoint. Charts are native PowerPoint XY charts, editable in place.

Writes Cluster_payload_stability.pptx. Run: python make_payload_ppt.py
(after draft_hydrostatics.py, cluster_statics.py, floatsim_crosscheck.py, payload_figures.py)
"""

# ruff: noqa: E501, E702, RUF001  -- slide-copy strings kept whole for editing; compact slide-
# assembly lines; degree signs, dashes and times signs in the slide text
from __future__ import annotations

import json
from pathlib import Path

import numpy as np
from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
HYD = json.loads((HERE / "buoy_hydrostatics_vs_draft.json").read_text(encoding="utf-8"))
ENV = json.loads((HERE / "cluster_envelope.json").read_text(encoding="utf-8"))
TILT = json.loads((HERE / "tilt_paths.json").read_text(encoding="utf-8"))
XCHK = json.loads((HERE / "floatsim_crosscheck.json").read_text(encoding="utf-8"))

NAVY, INK, MUTED = RGBColor(0x14, 0x23, 0x2E), RGBColor(0x2B, 0x3A, 0x44), RGBColor(0x5C, 0x6B, 0x75)
TEAL, WATER = RGBColor(0x0C, 0x8B, 0x96), RGBColor(0xE4, 0xF2, 0xF4)
AMBER, RED = RGBColor(0xC9, 0x77, 0x2B), RGBColor(0xB2, 0x43, 0x2C)
GREEN, WHITE = RGBColor(0x5A, 0x8A, 0x2B), RGBColor(0xFF, 0xFF, 0xFF)
ICE = RGBColor(0xB9, 0xD6, 0xDC)
DIR_COL = {0.0: TEAL, 15.0: GREEN, 30.0: AMBER, 45.0: RED}
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ helpers
def slide(dark: bool = False, notes: str = ""):  # type: ignore[no-untyped-def]
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY if dark else WHITE
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def text(s, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,  # type: ignore[no-untyped-def]
         anchor=MSO_ANCHOR.TOP, gap=6):
    """paras: a str, or a list of paragraphs; a paragraph is a str or a list of
    (text, {bold, color, size, italic}) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.0)
    for i, para in enumerate([paras] if isinstance(paras, str) else paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(gap)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run(); r.text = t
            r.font.name = FONT; r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold); r.font.italic = o.get("italic", False)
            r.font.color.rgb = o.get("color", color)
    return tb


def title(s, t, sub=None, dark=False):  # type: ignore[no-untyped-def]
    text(s, 0.6, 0.38, 12.1, 0.75, t, size=32, bold=True, color=WHITE if dark else NAVY)
    if sub:
        text(s, 0.6, 1.08, 12.1, 0.45, sub, size=16, color=ICE if dark else MUTED)


def card(s, x, y, w, h, fill=WATER):  # type: ignore[no-untyped-def]
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.06
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def stat(s, x, y, w, big, small, color=TEAL, fill=WATER, h=1.65):  # type: ignore[no-untyped-def]
    card(s, x, y, w, h, fill)
    text(s, x + 0.25, y + 0.15, w - 0.5, 0.75, big, size=34, bold=True, color=color)
    text(s, x + 0.25, y + 0.88, w - 0.5, h - 0.95, small, size=13, color=INK)


def picture(s, path, x, y, w=None, h=None):  # type: ignore[no-untyped-def]
    return s.shapes.add_picture(str(HERE / path), Inches(x), Inches(y),
                                Inches(w) if w else None, Inches(h) if h else None)


def fit_picture(s, path, x, y, w, h):  # type: ignore[no-untyped-def]
    """The picture scaled to fit inside the box (x, y, w, h), centred in it."""
    from PIL import Image
    iw, ih = Image.open(path if Path(path).is_absolute() else HERE / path).size
    k = min(w / iw, h / ih)
    pw, ph = iw * k, ih * k
    return s.shapes.add_picture(str(path if Path(path).is_absolute() else HERE / path),
                                Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw))


def xy_chart(s, x, y, w, h, series, xt, yt, xr=None, yr=None, legend=True,  # type: ignore[no-untyped-def]
             title_text=None):
    """series: list of dicts {name, x, y, color, width, dash, marker}."""
    cd = XyChartData()
    for sr in series:
        se = cd.add_series(sr["name"])
        for a, b in zip(sr["x"], sr["y"], strict=True):
            se.add_data_point(float(a), float(b))
    gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(x), Inches(y),
                            Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.name = FONT; ch.font.size = Pt(11); ch.font.color.rgb = MUTED
    ch.has_title = title_text is not None
    if title_text:
        ch.chart_title.text_frame.text = title_text
        tp = ch.chart_title.text_frame.paragraphs[0]
        tp.runs[0].font.size = Pt(13); tp.runs[0].font.bold = True; tp.runs[0].font.color.rgb = INK
    for sr, plot_s in zip(series, ch.plots[0].series, strict=True):
        ln = plot_s.format.line
        ln.color.rgb = sr["color"]; ln.width = Pt(sr.get("width", 2.25))
        if sr.get("dash"):
            ln.dash_style = sr["dash"]
        plot_s.smooth = False
        if sr.get("marker"):
            plot_s.marker.style = sr["marker"]; plot_s.marker.size = 6
            plot_s.marker.format.fill.solid(); plot_s.marker.format.fill.fore_color.rgb = sr["color"]
            plot_s.marker.format.line.color.rgb = sr["color"]
        else:
            plot_s.marker.style = XL_MARKER_STYLE.NONE
    for ax, t, rng in ((ch.category_axis, xt, xr), (ch.value_axis, yt, yr)):
        ax.has_title = True; ax.axis_title.text_frame.text = t
        r0 = ax.axis_title.text_frame.paragraphs[0].runs[0]
        r0.font.size = Pt(12); r0.font.bold = False; r0.font.color.rgb = INK
        ax.tick_labels.font.size = Pt(11); ax.tick_labels.font.color.rgb = MUTED
        ax.format.line.color.rgb = ICE
        ax.has_major_gridlines = True; ax.major_gridlines.format.line.color.rgb = RGBColor(0xE6, 0xEC, 0xEE)
        if rng:
            ax.minimum_scale, ax.maximum_scale = rng
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(11); ch.legend.font.color.rgb = INK
    return ch


def arrow(s, x, y, w=0.45, h=0.32, color=TEAL):  # type: ignore[no-untyped-def]
    sh = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color; sh.line.fill.background()


# ------------------------------------------------------------------ data
rows36 = [r for r in HYD["rows"]["36"] if r["sink_m"] >= -1e-9]
L_PIN = HYD["pin_above_cog_m"]
sink = [r["sink_m"] for r in rows36]
pin_load = [r["dW_N"] for r in rows36]
c55 = [r["C55"] for r in rows36]
destab = [p * L_PIN for p in pin_load]
net = [a - b for a, b in zip(c55, destab, strict=True)]
c33 = rows36[0]["C33"]
reserve = c33 * HYD["z_pin_design_m"]
m_centred = next(r["m_limit_kg"] for r in ENV["rows"] if r["offset_m"] == 0.0)
m_hub = ENV["m_hub_kg"]; m_cluster = 4 * ENV["m_buoy_kg"] + m_hub
k_tilt = TILT["k_tilt_Nm_per_rad"]


def me_at(tilt_deg: str) -> float:
    """mass x offset (kg m) for a hub tilt, from the 0.1 m branch (direction-independent)."""
    r = next(x for x in ENV["rows"] if x["dir_deg"] == 0.0 and x["offset_m"] == 0.1)
    return r["m_tilt_kg"][tilt_deg] * 0.1


def row(d: float, e: float) -> dict:
    return next(r for r in ENV["rows"] if r["dir_deg"] == d and abs(r["offset_m"] - e) < 1e-9)


G_ACC = HYD["g"]
sink_per_kg_mm = 1e3 * G_ACC / (4 * c33)
worst_pct = max(100 * (row(0.0, e)["m_limit_kg"] - row(45.0, e)["m_limit_kg"])
                / row(0.0, e)["m_limit_kg"] for e in np.arange(0.05, 0.501, 0.025))
gimbal_min = min(min(r["critical_state"]["gimbal_net_Nm_per_rad"]) for r in ENV["rows"])
roll_me = float(np.median([r["m_limit_kg"] * r["offset_m"] for r in ENV["rows"]
                           if r["limit_reason"].startswith("hub on its side")]))

# ================================================================== 1 title
s = slide(dark=True, notes=(
    "This deck answers one question: how much mass can sit on top of the 4-buoy cluster, and "
    "where, before its stability is compromised. Every number comes from the study's records; "
    "the drawings are to scale and come from the solved equilibria. The animations play in "
    "PowerPoint's slide show."))
text(s, 0.8, 1.6, 7.4, 0.5, "FLUME TEST ARTICLES · 4-BUOY CLUSTER", size=14, bold=True, color=ICE)
text(s, 0.8, 2.1, 7.4, 2.2, "How much payload can the cluster carry?", size=44, bold=True,
     color=WHITE)
text(s, 0.8, 4.35, 7.2, 1.2, "A mass on the hub, moved around: the stability envelope and the "
     "physics behind it", size=20, color=ICE)
text(s, 0.8, 6.5, 7.2, 0.4, "Payload-stability study · 2026-10-14 · Capytaine hydrostatics, "
     "large-angle statics, FloatSim cross-check", size=12, color=ICE)
card(s, 8.3, 0.6, 4.5, 6.3, WHITE)
fit_picture(s, "sketch_cluster.png", 8.45, 0.75, 4.2, 6.0)

# ================================================================== 2 bottom line
s = slide(notes=(
    "Four numbers carry the story. A centred mass only sinks the cluster: 54.7 kg puts all four "
    "pins at the waterline. Off-centre, the hub rocks like a seesaw on buoys that stay upright; "
    "every 0.6 kg·m of mass times offset gives about 5 degrees of hub tilt, in any direction. The "
    "buoys never flop over on their gimbals. And every limit is reached at 30 to 85 degrees of hub "
    "tilt, so in practice an allowable tilt is the better operating limit."))
title(s, "The bottom line", "Payload on the hub, free-floating in calm water, hub free to tilt, "
      "pins kept dry")
stat(s, 0.6, 1.9, 2.9, f"{m_centred:.1f} kg", "centred mass until all four pins reach the "
     "waterline: the cluster only sinks")
stat(s, 3.7, 1.9, 2.9, f"{me_at('5'):.1f} kg·m", "mass × offset that tilts the hub 5°, in every "
     "direction: the hub is a seesaw", color=AMBER)
stat(s, 6.8, 1.9, 2.9, "Never", f"do the buoys flop on their gimbals: margin ≥ {gimbal_min:.0f} "
     "N·m/rad at every load", color=GREEN)
stat(s, 9.9, 1.9, 2.85, "30–85°", "of hub tilt when an off-centre limit is reached: use a tilt "
     "limit to operate", color=RED)
text(s, 0.6, 4.0, 12.1, 0.5, "What limits the cluster", size=20, bold=True, color=NAVY)
text(s, 0.6, 4.55, 12.1, 2.4, [
    [("Centred: ", {"bold": True}), ("reserve buoyancy. ", {}),
     (f"Each kg sinks it {sink_per_kg_mm:.1f} mm; at {m_centred:.1f} kg the spar tops and pins are awash.", {})],
    [("Off-centre, up to ~0.3 m: ", {"bold": True}),
     ("the first pin goes awash, e.g. ", {}),
     (f"{row(45.0, 0.1)['m_limit_kg']:.1f} kg at 0.1 m or {row(45.0, 0.2)['m_limit_kg']:.1f} kg at "
      "0.2 m towards a buoy.", {})],
    [("Off-centre, beyond ~0.3 m: ", {"bold": True}),
     (f"the hub rolls onto its side at mass × offset ≈ {roll_me:.1f} kg·m; buoys may touch first.",
      {})],
    [("For operation: ", {"bold": True, "color": TEAL}),
     (f"≤ 10° of hub tilt ↔ mass × offset ≤ {me_at('10'):.1f} kg·m.", {"color": TEAL})],
], size=17, gap=8)

# ================================================================== 3 the cluster
s = slide(notes=(
    "The cluster is four OSU buoys hung from a small hub by gimbals at the spar tops, 0.717 m above "
    "the water. The gimbals leave roll and pitch free, so each buoy can stay vertical whatever the "
    "hub does. The study puts a point mass on the hub, in the pin plane, and moves it horizontally. "
    "The decisions: free-floating, calm water, the hub free to tilt, and the pins must stay dry."))
title(s, "The cluster and the question", "Four buoys hung from a hub on gimbals; a mass on the "
      "hub, moved parallel to the waterline")
fit_picture(s, "sketch_cluster.png", 0.6, 1.6, 6.0, 5.6)
text(s, 7.0, 1.75, 5.8, 0.5, "The question", size=20, bold=True, color=NAVY)
text(s, 7.0, 2.25, 5.8, 1.2, "How much mass can sit on the hub, and how far from the centre, "
     "before stability is compromised? Then: an envelope of mass and centre of gravity.", size=16)
text(s, 7.0, 3.6, 5.8, 0.5, "The set-up", size=20, bold=True, color=NAVY)
text(s, 7.0, 4.1, 5.8, 3.0, [
    [("Mass: ", {"bold": True}), ("a point mass on the hub, in the pin plane (+0.717 m), "
                                  "moved horizontally up to 0.5 m.", {})],
    [("Water: ", {"bold": True}), ("calm; the cluster free-floating (no mooring).", {})],
    [("Hub: ", {"bold": True}), ("free to tilt on the gimbals, with no mechanical stop.", {})],
    [("Limit: ", {"bold": True}), ("loss of restoring, or a pin reaching the water.", {})],
    [("Cluster: ", {"bold": True}), (f"{m_cluster:.1f} kg unloaded: 4 × {ENV['m_buoy_kg']} kg "
                                     f"buoys + {m_hub:.0f} kg hub.", {})],
], size=15, gap=7)

# ================================================================== 4 method
s = slide(notes=(
    "FloatSim can't reach these limits on its own: its restoring is a constant matrix fixed at one "
    "draft, it is only valid to about 5.7 degrees, and its linear model never loads the joints with "
    "the hub's weight. So the buoy's hydrostatics were computed with Capytaine at every draft, the "
    "cluster was solved exactly at any tilt with those hydrostatics, and FloatSim's own static solve "
    "was used to check the small-angle range, where the two agree exactly."))
title(s, "How it was computed", "FloatSim's restoring is fixed at one draft and valid to ~5.7°: "
      "the limits lie far beyond both")
steps = [("1  Capytaine at every draft",
          "The buoy re-meshed sunk −0.95 … +0.72 m: buoyancy, pin load and tilt stiffness at each "
          "draft (36 panels round; 96 reproduces the record's 74.08 N·m/rad).", TEAL),
         ("2  Large-angle statics",
          "Rigid hub on four buoys that hang vertically from their pins. The equilibrium at any tilt "
          "minimises the potential energy; its stiffness and each buoy's gimbal margin decide "
          "stability.", AMBER),
         ("3  FloatSim cross-check",
          "FloatSim's own cluster model, solved statically with its joints, against the statics at "
          "small tilt: identical sinkage, tilt within 0.002°.", GREEN)]
for i, (h, b, col) in enumerate(steps):
    x = 0.6 + i * 4.2
    card(s, x, 2.0, 3.7, 2.75)
    text(s, x + 0.25, 2.2, 3.2, 0.55, h, size=18, bold=True, color=col)
    text(s, x + 0.25, 2.85, 3.2, 2.5, b, size=14)
    if i < 2:
        arrow(s, x + 3.75, 3.2)
text(s, 0.6, 5.2, 12.1, 1.2, [
    [("Why not FloatSim alone: ", {"bold": True}),
     ("its restoring does not change as the cluster sinks (Level 2 is not wired), its kinematics "
      "are small-angle, and its linear model has no weight to load the joints. An approved "
      "exception to the 'use FloatSim' rule, with FloatSim checking the small-angle range.", {})]],
    size=14, color=MUTED)

# ================================================================== 5 physics 1: one buoy
s = slide(notes=(
    "What one buoy feels as the hub pushes it down. Left: the extra buoyancy grows in a straight "
    "line, 194 newtons per metre, because the spar's cross-section is constant, until the spar top "
    "goes under at 0.717 m; that caps each buoy at about 139 newtons, 56.8 kg for the four including "
    "the hub. Right: why the buoy doesn't flop over on its gimbal. The pin load at the top tries to "
    "tip it (red), but the buoy's own restoring grows faster as it sinks (blue), because the new "
    "buoyancy is added near the waterline. The margin (green) never drops below 24 N·m/rad."))
title(s, "What one buoy feels as it sinks", "The pin load grows linearly; the buoy's own "
      "restoring outgrows the tipping effect of that load")
xy_chart(s, 0.5, 1.6, 6.1, 4.6, [
    {"name": "pin load holding the sinkage (N)", "x": sink, "y": pin_load, "color": TEAL}],
    "sinkage below the design draft (m)", "pin load per buoy (N)", (0, 0.75), (0, 150), legend=False,
    title_text=f"Extra buoyancy: {c33:.0f} N/m until the spar top goes under")
xy_chart(s, 6.8, 1.6, 6.1, 4.6, [
    {"name": "buoy's restoring, C55", "x": sink, "y": c55, "color": RGBColor(0x2F, 0x74, 0xC0)},
    {"name": "tipping by the pin load, P × 1.62 m", "x": sink, "y": destab, "color": RED},
    {"name": "net margin", "x": sink, "y": net, "color": GREEN, "width": 3}],
    "sinkage below the design draft (m)", "stiffness about the buoy CoG (N·m/rad)", (0, 0.75),
    (0, 260), title_text="Gimbal stability of one buoy")
text(s, 0.6, 6.35, 12.1, 0.9, [
    [(f"Reserve: {reserve:.0f} N per buoy before its pin is awash → {m_centred:.1f} kg on the hub. ",
      {"bold": True}),
     (f"The net gimbal margin never drops below {gimbal_min:.0f} N·m/rad: the buoys stay upright "
      "at every load.", {})]], size=14)

# ================================================================== 6 scenario A
s = slide(notes=(
    "Scenario one: the mass in the centre. Nothing tilts; the four buoys share the load equally and "
    "the cluster sinks 12.6 mm per kilogram. The hub's own 2 kg already sinks it 25 mm. At 54.7 kg "
    "the spar tops, and the pins with them, reach the waterline. That is the centred limit: the "
    "cluster runs out of reserve buoyancy, it does not capsize."))
title(s, "A centred mass only sinks the cluster", "All four buoys share the load; nothing tilts")
fit_picture(s, "seq_centred.png", 0.6, 1.55, 12.1, 3.75)
stat(s, 0.6, 5.55, 3.8, f"{sink_per_kg_mm:.1f} mm", "of sinkage per kg of payload (194 N/m per buoy, four buoys)",
     h=1.5)
stat(s, 4.75, 5.55, 3.8, f"{m_centred:.1f} kg", "on the hub puts the spar tops and pins at the "
     "waterline", color=RED, h=1.5)
stat(s, 8.9, 5.55, 3.8, "0°", "hub tilt: a centred mass never tips the cluster", color=GREEN, h=1.5)

# ================================================================== 7 scenario B
r02 = row(45.0, 0.2)
s = slide(notes=(
    "Scenario two: the mass 0.2 m from the centre, towards one buoy. The hub rocks like a seesaw. "
    "The buoys stay vertical on their gimbals; the loaded buoy sinks, the far one barely moves, and "
    "only that difference in sinkage resists the tilt. That resistance is small, 67 N·m per radian, "
    "because the buoys are only 0.42 m from the centre. 3 kg already tilts the hub 5 degrees; at "
    f"{r02['m_limit_kg']:.1f} kg the loaded pin reaches the water with the hub at "
    f"{r02['tilt_at_limit_deg']:.0f} degrees."))
title(s, "Off-centre, the hub rocks like a seesaw", "0.2 m towards a buoy: the buoys stay upright; "
      "only their difference in sinkage resists the tilt")
fit_picture(s, "seq_offset.png", 0.6, 1.55, 12.1, 3.75)
stat(s, 0.6, 5.5, 3.8, f"{k_tilt:.0f} N·m/rad", "hub tilt stiffness: buoy sinkage × buoy "
     "distance² (0.42 m from the centre)", h=1.6)
stat(s, 4.75, 5.5, 3.8, "3 kg → 5°", "at 0.2 m: a few kilograms tilt the hub visibly",
     color=AMBER, h=1.6)
stat(s, 8.9, 5.5, 3.8, f"{r02['m_limit_kg']:.1f} kg", f"first pin awash, hub at "
     f"{r02['tilt_at_limit_deg']:.0f}° tilt", color=RED, h=1.6)

# ================================================================== 8 physics 2: m x e
lin_me = np.linspace(0.0, 7.0, 30)
series = []
for p in TILT["paths"]:
    series.append({"name": f"{p['offset_m']:.2f} m, {p['dir_deg']:.0f}°", "x": p["me_kgm"],
                   "y": p["tilt_deg"], "color": DIR_COL[p["dir_deg"]], "width": 1.75})
series.append({"name": "small-angle: θ = m g e / 67 N·m/rad", "x": lin_me,
               "y": np.degrees(lin_me * TILT["g"] / k_tilt), "color": INK, "width": 1.5,
               "dash": MSO_LINE_DASH_STYLE.DASH})
s = slide(notes=(
    "Why one number describes the tilt. Every branch, at six offsets and in two directions, falls on "
    "one curve when plotted against mass times offset: the payload's moment is what tilts the hub, "
    "and the hub's resistance is the same in every direction because the square of buoys is "
    "symmetric. For small tilts the curve follows theta equals m g e over 67 N·m per radian (dashed); "
    "at larger tilts it bends because the lever arms shorten as the hub turns."))
title(s, "The tilt depends only on mass × offset", "Every branch, in every direction, falls on "
      "one curve")
xy_chart(s, 0.5, 1.55, 7.6, 5.6, series, "mass × offset (kg·m)", "hub tilt (°)", (0, 7), (0, 90),
         legend=False, title_text="Hub tilt against payload moment: 12 branches")
text(s, 8.4, 1.6, 4.4, 0.5, "Reading it", size=20, bold=True, color=NAVY)
tb = s.shapes.add_table(6, 3, Inches(8.4), Inches(2.2), Inches(4.4), Inches(2.6)).table
for j, hdr in enumerate(("Hub tilt", "mass × offset", "e.g. at 0.2 m")):
    tb.cell(0, j).text = hdr
for i, t in enumerate(("2", "5", "10", "20", "30"), start=1):
    me = me_at(t)
    for j, v in enumerate((f"{t}°", f"{me:.2f} kg·m", f"{me / 0.2:.1f} kg")):
        tb.cell(i, j).text = v
for i in range(6):
    for j in range(3):
        c = tb.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = TEAL if i == 0 else (
            WATER if i % 2 else WHITE)
        for p in c.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(14); r.font.name = FONT
                r.font.color.rgb = WHITE if i == 0 else INK; r.font.bold = i == 0
text(s, 8.4, 5.1, 4.4, 2.0, [
    [("Lines: ", {"bold": True}), ("12 branches, offsets 0.05–0.5 m, between buoys (teal) and "
                                  "towards a buoy (red): they lie on top of each other, each ending "
                                  "at its own limit. Dashed: the small-angle formula.", {})]],
    size=13, color=MUTED)

# ================================================================== 9 scenario C
r04 = row(0.0, 0.4)
s = slide(notes=(
    "Scenario three: the mass 0.4 m out, between two buoys, near the edge of the cluster. Here the "
    "hub keeps tilting long before any pin reaches the water, and rolls onto its side: beyond about "
    "85 degrees the payload is no longer on top. That happens at a mass times offset of about "
    f"{roll_me:.1f} kg·m in every direction. Before that, as the hub turns, the hanging buoys swing "
    "towards each other and may touch."))
title(s, "Far out, the hub rolls onto its side", "0.4 m between buoys: no pin reaches the water "
      "before the hub turns over")
fit_picture(s, "seq_rollover.png", 0.6, 1.55, 12.1, 3.75)
stat(s, 0.6, 5.5, 3.8, f"{roll_me:.1f} kg·m", "mass × offset at which the hub rolls onto its side, "
     "in every direction", color=RED, h=1.6)
stat(s, 4.75, 5.5, 3.8, f"{r04['m_limit_kg']:.1f} kg", f"at 0.4 m; the lowest pin is still "
     f"{1e3 * r04['critical_state']['min_pin_freeboard_m']:.0f} mm above water", color=AMBER, h=1.6)
stat(s, 8.9, 5.5, 3.8, f"{r04['m_buoys_touch_kg']:.1f} kg", "at 0.4 m the buoys may already touch "
     "each other", color=MUTED, h=1.6)

# ================================================================== 10 animations
s = slide(notes=(
    "The three loadings animated, the mass rising from zero to the limit. They play in slide show "
    "(and in normal view in recent PowerPoint). Watch the centred cluster sink straight down, the "
    "seesaw tilt with the buoys hanging vertically, and the hub roll over when the mass is far out."))
title(s, "Watch it happen", "The mass rising from zero to the limit (animations play in slide "
      "show)")
for i, (g, lab) in enumerate((("anim_centred.gif", "Centred"),
                              ("anim_offset.gif", "0.2 m towards a buoy"),
                              ("anim_rollover.gif", "0.4 m between buoys"))):
    x = 0.55 + i * 4.15
    picture(s, g, x, 1.65, w=4.0)
    text(s, x, 5.5, 4.0, 0.4, lab, size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

# ================================================================== 11 envelope
env_series = []
for d in (0.0, 15.0, 30.0, 45.0):
    rr = [r for r in ENV["rows"] if r["dir_deg"] == d]
    lab = {0.0: "0° (between buoys)", 45.0: "45° (towards a buoy)"}.get(d, f"{d:.0f}°")
    env_series.append({"name": f"limit, {lab}", "x": [r["offset_m"] for r in rr],
                       "y": [r["m_limit_kg"] for r in rr], "color": DIR_COL[d], "width": 2.5})
r0 = [r for r in ENV["rows"] if r["dir_deg"] == 0.0 and r["offset_m"] > 0]
TILT_SHADE = {"5": INK, "10": MUTED, "20": RGBColor(0xA8, 0xB4, 0xBA)}
for t, w_ in (("5", 1.75), ("10", 1.5), ("20", 1.5)):
    pts = [(r["offset_m"], r["m_tilt_kg"][t]) for r in r0 if r["m_tilt_kg"][t] is not None]
    env_series.append({"name": f"hub tilt {t}°", "x": [a for a, _ in pts],
                       "y": [b for _, b in pts], "color": TILT_SHADE[t], "width": w_,
                       "dash": MSO_LINE_DASH_STYLE.DASH})
touch = []
for r in r0:
    vals = [x["m_buoys_touch_kg"] for x in ENV["rows"]
            if x["offset_m"] == r["offset_m"] and x["m_buoys_touch_kg"]]
    if vals:
        touch.append((r["offset_m"], min(vals)))
env_series.append({"name": "buoys may touch (worst direction)", "x": [a for a, _ in touch],
                   "y": [b for _, b in touch], "color": RGBColor(0x7B, 0x3F, 0xA0), "width": 1.75,
                   "dash": MSO_LINE_DASH_STYLE.ROUND_DOT})
s = slide(notes=(
    "The envelope. Each solid line is the payload limit for one direction of the offset: pins awash "
    "at small offsets, the hub rolling over beyond about 0.3 m. Towards a buoy (red) is the worst "
    "direction, by up to 15 percent. The dashed lines are hub tilts of 5, 10 and 20 degrees: they "
    "lie far inside the limits, which is why a tilt limit is the practical operating envelope. The "
    "dotted line is where the buoys may touch, which comes before the limit at large offsets."))
title(s, "The payload envelope", "Payload mass against its offset from the hub centre: limits "
      "(solid) and hub tilt (dashed)")
xy_chart(s, 0.5, 1.5, 8.6, 5.75, env_series, "payload offset from the hub centre (m)",
         "payload mass on the hub (kg)", (0, 0.5), (0, 60), title_text=None)
text(s, 9.4, 1.6, 3.4, 0.5, "Reading it", size=20, bold=True, color=NAVY)
text(s, 9.4, 2.15, 3.4, 5.0, [
    [("Below a solid line: ", {"bold": True}), ("stable and pins dry, for that direction.", {})],
    [("Worst direction: ", {"bold": True}), (f"towards a buoy (red), by up to {worst_pct:.0f} %.", {})],
    [("Up to ~0.3 m: ", {"bold": True}), ("a pin goes awash. ", {}),
     ("Beyond: ", {"bold": True}), ("the hub rolls over.", {})],
    [("Dashed: ", {"bold": True}), ("hub tilt 5° (darkest), 10°, 20° (lightest), far inside the "
                                   "limits.", {})],
    [("Dotted: ", {"bold": True}), ("buoys may touch, before the limit beyond ~0.2 m.", {})],
], size=14, gap=10)

# ================================================================== 12 mass-CG envelope
cg_series = []
for d in (0.0, 45.0):
    rr = [r for r in ENV["rows"] if r["dir_deg"] == d]
    lab = {0.0: "0° (between buoys)", 45.0: "45° (towards a buoy)"}[d]
    cg_series.append({"name": f"limit, {lab}",
                      "x": [1e3 * r["cg_at_limit"]["cg_offset_m"] for r in rr],
                      "y": [r["cg_at_limit"]["total_mass_kg"] for r in rr], "color": DIR_COL[d],
                      "width": 2.5})
for t in ("5", "10", "20"):
    pts = [(r["offset_m"], r["m_tilt_kg"][t]) for r in r0 if r["m_tilt_kg"][t] is not None]
    cg_series.append({"name": f"hub tilt {t}°",
                      "x": [1e3 * m * e / (m_cluster + m) for e, m in pts],
                      "y": [m_cluster + m for _, m in pts], "color": TILT_SHADE[t], "width": 1.5,
                      "dash": MSO_LINE_DASH_STYLE.DASH})
s = slide(notes=(
    "The same envelope expressed as the whole cluster's mass and centre of gravity, which is how a "
    "loading envelope is usually given. The cluster weighs 88 kg unloaded. Moving its overall "
    "centre of gravity just a few millimetres off-centre tilts the hub: about 7 mm for 5 degrees "
    "with a light payload. The limits need tens of millimetres."))
title(s, "Mass and CG envelope of the whole cluster", f"Total mass against the horizontal offset "
      f"of the cluster's CG (unloaded {m_cluster:.1f} kg, level geometry)")
xy_chart(s, 0.5, 1.5, 8.6, 5.75, cg_series, "horizontal offset of the cluster's CG (mm)",
         "total cluster mass (kg)", (0, 70), (85, 145))
text(s, 9.4, 1.6, 3.4, 0.5, "At the limit", size=20, bold=True, color=NAVY)
tb = s.shapes.add_table(6, 3, Inches(9.4), Inches(2.2), Inches(3.4), Inches(3.0)).table
for j, hdr in enumerate(("Payload at", "Total mass", "CG offset")):
    tb.cell(0, j).text = hdr
for i, e in enumerate((0.0, 0.1, 0.2, 0.3, 0.4), start=1):
    r = row(45.0, e)
    for j, v in enumerate((f"{e:.1f} m", f"{r['cg_at_limit']['total_mass_kg']:.1f} kg",
                           f"{1e3 * r['cg_at_limit']['cg_offset_m']:.0f} mm")):
        tb.cell(i, j).text = v
for i in range(6):
    for j in range(3):
        c = tb.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = TEAL if i == 0 else (
            WATER if i % 2 else WHITE)
        for p in c.text_frame.paragraphs:
            for rr_ in p.runs:
                rr_.font.size = Pt(13); rr_.font.name = FONT
                rr_.font.color.rgb = WHITE if i == 0 else INK; rr_.font.bold = i == 0
text(s, 9.4, 5.45, 3.4, 1.6, "Worst direction (towards a buoy). The CG offset is that of the "
     "whole cluster with the payload, before it tilts.", size=12, color=MUTED)

# ================================================================== 13 validation
s = slide(notes=(
    "How far to trust it. The hydrostatics reproduce the recorded pitch stiffness of the buoy on the "
    "fine mesh. And FloatSim, solving its own model of the cluster with its joints, gives the same "
    "sinkage and the same tilt as the large-angle statics in the range where FloatSim is valid. The "
    "large-angle part rests on simple, exact geometry: vertical buoys of constant section."))
title(s, "Cross-checks", "Capytaine against the record; FloatSim against the large-angle statics "
      "at small tilt")
stat(s, 0.6, 1.8, 3.6, f"{HYD['nt96_design']['C55']:.2f}", "N·m/rad: the buoy's pitch stiffness "
     "on the 96-panel mesh = the record (DESIGN-BASIS §C8)", color=GREEN, h=1.9)
stat(s, 0.6, 3.95, 3.6, "0.002°", "largest tilt difference between FloatSim and the statics "
     "(2 kg at 0.2 m)", color=GREEN, h=1.9)
n = len(XCHK) + 1
tb = s.shapes.add_table(n, 5, Inches(4.6), Inches(1.8), Inches(8.1), Inches(0.45 * n)).table
for j, hdr in enumerate(("Payload", "Offset, direction", "FloatSim sinkage / tilt",
                         "Statics sinkage / tilt", "Agree")):
    tb.cell(0, j).text = hdr
for i, x in enumerate(XCHK, start=1):
    f, st = x["floatsim"], x["statics"]
    vals = (f"{x['m_kg']:.0f} kg", f"{x['offset_m']:.2f} m, {x['dir_deg']:.0f}°",
            f"{-1e3 * f['heave_m']:.1f} mm / {f['tilt_deg']:.3f}°",
            f"{-1e3 * st['heave_m']:.1f} mm / {st['tilt_deg']:.3f}°",
            "yes" if abs(f["tilt_deg"] - st["tilt_deg"]) < 0.01 else "check")
    for j, v in enumerate(vals):
        tb.cell(i, j).text = v
for i in range(n):
    for j in range(5):
        c = tb.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = TEAL if i == 0 else (
            WATER if i % 2 else WHITE)
        for p in c.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(13); r.font.name = FONT
                r.font.color.rgb = WHITE if i == 0 else INK; r.font.bold = i == 0
text(s, 4.6, 1.95 + 0.45 * n, 8.1, 1.0, "Both on the same 12-panel hydrostatics; sinkage is "
     "measured from the unloaded draft and includes the hub's 2 kg.", size=12, color=MUTED)

# ================================================================== 14 what it means
s = slide(notes=(
    "What to take away. Keep the payload central: the cluster can then carry tens of kilograms. "
    "Off-centre, plan by hub tilt, not by the stability limit, which only comes at extreme tilts. "
    "The analysis is static and in calm water; waves, the mooring, a raised payload and the "
    "gimbals' real travel would all reduce the margins."))
title(s, "What it means for loading the cluster")
card(s, 0.6, 1.5, 5.95, 4.3)
text(s, 0.85, 1.7, 5.5, 0.5, "Guidance", size=20, bold=True, color=TEAL)
text(s, 0.85, 2.3, 5.5, 4.5, [
    [("Keep the mass central. ", {"bold": True}), ("Centred, the cluster carries "
     f"{m_centred:.1f} kg before its pins are awash; it only sinks, {sink_per_kg_mm:.1f} mm per kg.", {})],
    [("Plan off-centre loads by hub tilt. ", {"bold": True}),
     (f"5° ↔ {me_at('5'):.1f} kg·m, 10° ↔ {me_at('10'):.1f} kg·m, in every direction.", {})],
    [("Worst direction: ", {"bold": True}), (f"towards a buoy, by up to {worst_pct:.0f} % on the limit.", {})],
    [("Don't rely on the limits for operation: ", {"bold": True}),
     ("they come at 30–85° of hub tilt, where payload retention, the gimbals and buoy contact "
      "govern.", {})]], size=15, gap=10)
card(s, 6.8, 1.5, 5.95, 4.3, RGBColor(0xF6, 0xEE, 0xE8))
text(s, 7.05, 1.7, 5.5, 0.5, "Not included", size=20, bold=True, color=RED)
text(s, 7.05, 2.3, 5.5, 4.5, [
    [("Waves and dynamics: ", {"bold": True}), ("this is calm-water statics; waves reduce the "
                                                "margins.", {})],
    [("The mooring: ", {"bold": True}), ("free-floating; the moored cluster's 1° pretension lean "
                                         "is not included.", {})],
    [("A raised payload: ", {"bold": True}), ("a mass above the pin plane adds a tipping moment "
                                               "to the hub; needs a new run.", {})],
    [("Hardware: ", {"bold": True}), ("gimbal travel, payload retention and hub buoyancy are not "
                                      "modelled; buoy contact is only flagged.", {})]],
    size=15, gap=10)

# ================================================================== 15 next
s = slide(dark=True, notes=(
    "Next, the same exercise for the 2.4 m square platform. Its deck sits at 0.90 m, above the "
    "pins, so a mass on the deck tips the hubs as they tilt: the term that is zero for the cluster. "
    "The statics need that extension; the Capytaine draft table is reused as it is."))
text(s, 0.8, 1.2, 11.5, 0.5, "NEXT", size=14, bold=True, color=ICE)
text(s, 0.8, 1.7, 11.5, 1.2, "The same envelope for the 2.4 m platform", size=40, bold=True,
     color=WHITE)
card(s, 7.3, 2.95, 5.5, 3.6, WHITE)
fit_picture(s, HERE.parent / "flume-mooring" / "figs" / "plan_platform.png", 7.45, 3.1, 5.2, 3.3)
text(s, 7.3, 6.65, 5.5, 0.4, "The 2.4 m platform in the flume (mooring rev D), to scale",
     size=12, color=ICE)
text(s, 0.8, 3.0, 6.2, 3.6, [
    [("Same method: ", {"bold": True, "color": WHITE}),
     ("the Capytaine draft table is reused; the statics extend to four hubs and the deck.",
      {"color": ICE})],
    [("New physics: ", {"bold": True, "color": WHITE}),
     ("the deck sits at +0.90 m, above the pins, so a mass on it tips the hubs as they tilt, a "
      "term that is zero for the cluster.", {"color": ICE})],
    [("Then: ", {"bold": True, "color": WHITE}),
     ("a check in waves for the loads that matter, if needed.", {"color": ICE})]],
    size=17, gap=12)

out = HERE / "Cluster_payload_stability.pptx"
prs.save(str(out))
print(f"wrote {out.name} ({len(prs.slides)} slides)")
