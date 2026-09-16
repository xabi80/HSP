"""Render REBUTTAL-sidewall.md into a submission-ready Word document (with the two key
figures embedded) for the TEAMER response. Requires python-docx (scripting-only tool).

Run: ``python build_rebuttal_docx.py`` -> REBUTTAL-sidewall.docx next to this script.
"""
# ruff: noqa: RUF001  -- the rendered DOCX text intentionally uses typographic
# minus/en-dash/<= glyphs; RUF001 flags these string chars as "ambiguous unicode".
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
TEAL, INK, GREY = RGBColor(0x0C, 0x8B, 0x96), RGBColor(0x25, 0x32, 0x3A), RGBColor(0x54, 0x63, 0x6D)


def _shade(el, fill: str) -> None:
    pr = el.get_or_add_tcPr() if el.tag.endswith("}tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _left_border(par, color: str) -> None:
    pPr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "10"), ("w:color", color)):
        left.set(qn(k), v)
    bdr.append(left)
    pPr.append(bdr)


def build() -> None:
    doc = Document()
    normal = doc.styles["Normal"].font
    normal.name, normal.size = "Calibri", Pt(10.5)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for m in ("top", "bottom"):
        setattr(sec, f"{m}_margin", Inches(0.75))
    sec.left_margin = sec.right_margin = Inches(0.85)

    def para(runs, after=8, align=None, style=None, before=0):  # type: ignore[no-untyped-def]
        p = doc.add_paragraph(style=style)
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        for text, o in runs:
            r = p.add_run(text)
            r.bold = o.get("b", False)
            r.italic = o.get("i", False)
            r.font.color.rgb = o.get("c", INK)
            r.font.size = Pt(o.get("size", 10.5))
        return p

    def heading(text):  # type: ignore[no-untyped-def]
        para([(text, {"b": True, "c": TEAL, "size": 13})], after=4, before=12)

    def bullet(runs):  # type: ignore[no-untyped-def]
        para(runs, after=4, style="List Bullet")

    def figure(fname, width_in, caption):  # type: ignore[no-untyped-def]
        doc.add_picture(str(HERE / fname), width=Inches(width_in))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        para([(caption, {"i": True, "c": GREY, "size": 8.5})], after=10,
             align=WD_ALIGN_PARAGRAPH.CENTER)

    # --- title ---
    para([("Response to reviewer comment", {"b": True, "c": INK, "size": 15})], after=2)
    para([("Sidewall effects on the Phase-3 platform test — HWRL Large Wave Flume",
           {"b": True, "c": TEAL, "size": 12})], after=10)

    # --- reviewer quote ---
    para([("Reviewer comment.", {"b": True})], after=4)
    q = para([("“Phase 3 platform (2.50 m) inside HWRL's 3.67 m flume leaves 0.6 m side "
               "clearance, inducing severe sidewall reflections that will corrupt dynamic data.”",
               {"i": True})], after=12)
    q.paragraph_format.left_indent = Inches(0.25)
    _shade(q._p, "F2F7F8")
    _left_border(q, "0C8B96")

    heading("Summary")
    para([("We assessed the concern quantitatively with potential-flow boundary-element "
           "simulations (Capytaine) in which the flume side walls are modelled ", {}),
          ("explicitly", {"i": True}),
          (". The result: ", {}), ("sidewall reflections do not corrupt the Phase-3 dynamic data.",
                                    {"b": True}),
          (" Modelling the walls in versus out (at the same 2.7 m operating depth) shifts every "
           "measured natural period by ", {}), ("< 0.7%", {"b": True}),
          (" and every wave-frequency load by ", {}), ("≤ 9%", {"b": True}),
          (" (≤ 2% for pitch and surge) — inside normal model-test uncertainty — with "
           "no spurious resonances in the operating band. The physical reason is that the platform "
           "is a very weak wavemaker in the modes being measured, so there is little radiated wave "
           "energy for the walls to reflect.", {})])

    heading("Why the standard “sidewall reflection” concern does not apply here")
    para([("“Severe reflections corrupt data” is the correct concern for bodies that are "
           "strong wavemakers or that span the flume. The Phase-3 platform is neither:", {})],
         after=4)
    bullet([("It barely radiates. ", {"b": True}),
            ("Its heave radiation damping is only 3.5% of the total heave damping (the rest is "
             "viscous); pitch and roll radiate far less, because the buoys move out of phase and "
             "their radiated waves largely cancel. With so little wave energy leaving the "
             "platform, there is almost nothing to reflect.", {})])
    bullet([("The operating band is sub-cutoff. ", {"b": True}),
            ("Reflections only build a coherent cross-flume standing wave at the flume's "
             "transverse cut-on periods, T ≈ 2.19 / 1.53 / 1.25 s. The heave resonance "
             "(≈ 2.5 s) is below the first cut-on, where the transverse field is evanescent "
             "and cannot organise into a standing wave.", {})])
    bullet([("It is porous and does not span the flume. ", {"b": True}),
            ("Water passes between the 12 buoys and escapes freely along the 104 m flume length; "
             "only the cross-flume direction is bounded, several plate-radii from each outer "
             "buoy.", {})])

    heading("Method")
    para([("Capytaine 2.3.1 potential-flow BEM. Platform = 12 spar-plate buoys at the Phase-3 "
           "layout (buoy centres on a 2.5 m circle), each buoy the Phase-1 decay-correlated hull "
           "(0.159 m spar, equal-area heave plate r = 0.144 m) that reproduced the Phase-1 "
           "free-decay (T ≈ 2.5 s, ζ ≈ 13%). Side walls (W = 3.66 m) imposed by the "
           "method of images; finite depth h = 2.7 m native. All comparisons are walls-in vs "
           "walls-out at the same depth, isolating the sidewall effect.", {})])
    figure("flume_blockage.png", 3.7,
           "Figure 1. The 12-buoy platform to scale in the 3.66 m flume "
           "(2.5 m to buoy centres; ~44 cm/side).")

    heading("Findings")
    para([("Free-decay tests (natural periods and damping):", {"b": True})], after=6, before=4)
    rows = [("Mode", "Wall shift in natural period", "Wall shift in added mass"),
            ("Heave", "−0.25%", "−1.5%"),
            ("Pitch (= roll)", "< 0.2%", "< 0.3%"),
            ("Surge (= sway)", "< 0.7%", "≤ 2%")]
    tbl = doc.add_table(rows=4, cols=3)
    tbl.style = "Table Grid"
    widths = (Inches(1.7), Inches(2.7), Inches(2.5))
    for i, (a, b, c) in enumerate(rows):
        cells = tbl.rows[i].cells
        for j, txt in enumerate((a, b, c)):
            cells[j].width = widths[j]
            cp = cells[j].paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = cp.add_run(txt)
            run.font.size = Pt(10)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                _shade(cells[j]._tc, "0C8B96")
            else:
                run.bold = j == 1
                run.font.color.rgb = INK
    para([("All shifts are well inside the ±(3–5)% uncertainty of a physical decay test. "
           "Because heave radiation is only 3.5% of the damping, even a hypothetical perfect wall "
           "reflection could perturb the measured damping by at most a few percent — and the "
           "explicit walls-in simulation shows less.", {})], before=8)
    para([("Wave-sweep tests (RAOs):", {"b": True})], after=6, before=4)
    bullet([("Sidewall effect on wave-frequency loads: ≤ 9% (heave, at resonance), "
             "≤ 2% (pitch, surge).", {})])
    bullet([("Heave RAO peak: 0.394 (walls out) → 0.359 (walls in), −9% — a smooth, "
             "predictable offset, not “corruption.”", {})])
    bullet([("Even at the transverse cut-ons the wall effect on the loads stays within ~±6% "
             "(refined to three image reflections); the platform's weak scattering prevents a "
             "sharp resonance.", {})])
    figure("rebuttal_summary.png", 6.6,
           "Figure 2. Walls-in vs walls-out at 2.7 m: free-decay period shifts (left) and "
           "wave-sweep load shifts (right), all modes.")

    heading("What we concede, and how we handle it")
    bullet([("Transverse cut-ons (T ≈ 2.19 / 1.53 / 1.25 s): ", {"b": True}),
            ("bounded to ~±6% even there, but as good practice we avoid parking sweep periods "
             "exactly on them — narrow, known, and easily skipped in the test matrix.", {})])
    bullet([("Clearance number: ", {"b": True}),
            ("the reviewer's 0.6 m assumes 2.50 m is the outer extent. Our 2.50 m is the "
             "buoy-centre circle; with the heave plates the outer clearance is ~0.44 m. We "
             "analysed the tighter, as-built ~0.44 m and the conclusions already reflect it.", {})])
    bullet([("Finite depth: ", {"b": True}),
            ("the separate, larger effect at long waves is the 2.7 m depth (not the walls) — a "
             "known facility property we account for when relating flume RAOs to the target "
             "environment.", {})])

    heading("Conclusion")
    para([("Explicit BEM modelling of the flume side walls shows the Phase-3 platform's measured "
           "natural periods change by < 0.7% and its wave-frequency loads by ≤ 9%, with no "
           "in-band spurious resonances. The platform's very low wave radiation and the sub-cutoff "
           "operating band mean sidewall reflections are negligible for both the free-decay and "
           "wave-sweep campaigns. The 2.50 m platform is compatible with the HWRL Large Wave "
           "Flume.", {})])

    out = HERE / "REBUTTAL-sidewall.docx"
    doc.save(str(out))
    print(f"wrote {out}")


if __name__ == "__main__":
    build()
